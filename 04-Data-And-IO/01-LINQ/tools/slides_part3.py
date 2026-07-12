from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from tools.builder_core import (
    add_bg, add_header, add_bottom_banner, add_code_panel,
    add_theory_card, add_dual_comparison_card, add_table_card,
    ACCENT_INDIGO, ACCENT_SKY, ACCENT_EMERALD, ACCENT_VIOLET,
    ACCENT_AMBER, ACCENT_TEAL, ACCENT_CRIMSON
)

def build_part3(prs, blank_layout):
    print("Building Part 3 (Slides 73 - 108)...")

    # ==========================================
    # SLIDE 73: EQUALITY OPERATOR: SEQUENCEEQUAL
    # ==========================================
    s73 = prs.slides.add_slide(blank_layout)
    add_bg(s73)
    add_header(s73, "12 · Порівняння (Equality)", "Оператор рівності послідовностей: SequenceEqual", 73)
    add_theory_card(s73, 0.8, 1.55, 5.72, 4.25,
        "Перевірка ідентичності двох колекцій", ACCENT_INDIGO,
        [
            ("Єдиний оператор рівності:", "У LINQ існує тільки один оператор рівності — SequenceEqual."),
            ("Три критерії збігу:", "Метод перевіряє: 1) однакову кількість елементів, 2) попарну рівність значень, 3) абсолютно однаковий порядок слідування."),
            ("Порівняння примітивів:", "Для простих типів (string, int) використовується стандартний компаратор EqualityComparer<T>.Default."),
            ("Порядок має значення:", "Навіть якщо колекції містять абсолютно ті самі елементи, але в іншому порядку — повертається false.")
        ],
        "⚡ Поведінка методу:",
        "Обхід обох колекцій відбувається синхронно через два ітератори MoveNext() паралельно."
    )
    add_code_panel(s73, 6.81, 1.55, 5.72, 4.25, [
        "// Приклад 1: Однаковий вміст і однаковий порядок",
        "IList<string> strList1 = new List<string>() {",
        "    \"One\", \"Two\", \"Three\", \"Four\", \"Three\"",
        "};",
        "IList<string> strList2 = new List<string>() {",
        "    \"One\", \"Two\", \"Three\", \"Four\", \"Three\"",
        "};",
        "bool isEqual1 = strList1.SequenceEqual(strList2);",
        "Console.WriteLine(isEqual1); // true",
        "",
        "// Приклад 2: Однаковий вміст, але різний порядок",
        "IList<string> strList3 = new List<string>() { \"One\", \"Two\", \"Three\" };",
        "IList<string> strList4 = new List<string>() { \"Two\", \"One\", \"Three\" };",
        "bool isEqual2 = strList3.SequenceEqual(strList4);",
        "Console.WriteLine(isEqual2); // false!",
        "",
        "-> Output: True",
        "-> Output: False"
    ], "Program.cs")
    add_bottom_banner(s73, "💡 Важливе правило:",
                      "SequenceEqual повертає true виключно за умови повного попарного збігу елементів та їхніх індексів у послідовності.")

    # ==========================================
    # SLIDE 74: SEQUENCEEQUAL ДЛЯ ОБ'ЄКТІВ ТА COMPARER
    # ==========================================
    s74 = prs.slides.add_slide(blank_layout)
    add_bg(s74)
    add_header(s74, "12 · Порівняння (Equality)", "Порівняння об'єктів Student за допомогою SequenceEqual та IEqualityComparer", 74)
    add_theory_card(s74, 0.8, 1.55, 5.72, 4.25,
        'Посилальні типи та кастомний компаратор', ACCENT_INDIGO,
        [
            ('Проблема посилань:', 'За замовчуванням для класів C# порівнює посилання в купі (ReferenceEquals). Два екземпляри з однаковими полями вважаються різними!'),
            ('Кастомний компаратор:', 'Щоб порівнювати екземпляри Student за бізнес-полями (StudentID, StudentName), передається IEqualityComparer<Student>.'),
            ('Реалізація методів:', 'Компаратор повинен реалізувати Equals(Student x, Student y) та узгоджений GetHashCode(Student obj).'),
            ('Вимога узгодженості:', "Якщо Equals повертає true для двох об'єктів, їхні GetHashCode обов'язково повинні збігатися для коректного хешування.")
        ],
        '⚡ Результат порівняння:',
        "За наявності StudentComparer метод повертає true, оскільки всі відповідні студенти мають однаковий ID та ім'я."
    )
    add_code_panel(s74, 6.81, 1.55, 5.72, 4.25, [
        'IList<Student> studentList1 = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 }',
        '};',
        '',
        'IList<Student> studentList2 = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 }',
        '};',
        '',
        "// Без компаратора поверне false (різні екземпляри в пам'яті)",
        '// З кастомним компаратором порівнюються поля сутності:',
        'bool isEqual = studentList1.SequenceEqual(',
        '    studentList2, new StudentComparer());',
        '',
        'Console.WriteLine(isEqual); // True',
        '-> Output: True'
    ], 'StudentComparer.cs')
    add_bottom_banner(s74, "📌 Угода про хешування:",
                      "Метод GetHashCode компаратора зобов'язаний повертати однаковий код для об'єктів, які метод Equals вважає рівними.")

    # ==========================================
    # SLIDE 75: CONCAT РЯДКОВИХ КОЛЕКЦІЙ
    # ==========================================
    s75 = prs.slides.add_slide(blank_layout)
    add_bg(s75)
    add_header(s75, "13 · Конкатенація (Concatenation)", "Оператор Concat: Послідовне злиття двох списків рядків", 75)
    add_theory_card(s75, 0.8, 1.55, 5.72, 4.25,
        'Послідовне приєднання колекцій', ACCENT_INDIGO,
        [
            ('Принцип Concat:', 'Метод Concat приєднує другу послідовність у кінець першої, створюючи єдиний потік даних.'),
            ('Збереження дублікатів:', 'На відміну від Union, оператор Concat НЕ усуває повторювані елементи і не змінює порядок.'),
            ('Швидкодія O(N + M):', "Метод не будує хеш-таблиць у пам'яті, тому працює надзвичайно швидко з мінімальними накладними витратами."),
            ('Потокова передача:', "Concat транслюється у два послідовних цикли foreach, не виділяючи проміжних буферних масивів у пам'яті.")
        ],
        '⚡ Результат виклику:',
        'Створюється послідовність: One, Two, Three, Five, Six. Елементи обходяться ліниво один за одним.'
    )
    add_code_panel(s75, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> collection1 = new List<string>() {',
        '    "One", "Two", "Three"',
        '};',
        '',
        'IList<string> collection2 = new List<string>() {',
        '    "Five", "Six"',
        '};',
        '',
        '// Послідовне злиття двох списків',
        'var collection3 = collection1.Concat(collection2);',
        '',
        'foreach (string str in collection3)',
        '    Console.WriteLine(str);',
        '',
        '-> Output:',
        '-> One',
        '-> Two',
        '-> Three',
        '-> Five',
        '-> Six'
    ], 'Program.cs')
    add_bottom_banner(s75, "💡 Відмінність Concat від Union:",
                      "Concat просто конкатенує послідовності без перевірки на унікальність. Якщо потрібна множина без повторів — використовуйте Union.")

    # ==========================================
    # SLIDE 76: CONCAT ЧИСЛОВИХ КОЛЕКЦІЙ
    # ==========================================
    s76 = prs.slides.add_slide(blank_layout)
    add_bg(s76)
    add_header(s76, "13 · Конкатенація (Concatenation)", "Оператор Concat: Злиття числових послідовностей int", 76)
    add_theory_card(s76, 0.8, 1.55, 5.72, 4.25,
        'Конкатенація примітивних значень', ACCENT_INDIGO,
        [
            ("Об'єднання чисел:", 'Демонстрація злиття двох числових послідовностей: [1, 2, 3] та [4, 5, 6].'),
            ('Лінивий конвеєр:', 'Використовує механізм yield return: спочатку повністю читається перша колекція, а потім — друга.'),
            ('Ланцюговий виклик:', "Можна об'єднувати декілька джерел: col1.Concat(col2).Concat(col3)."),
            ('Типізація результату:', 'Повертає IEnumerable<int>, зберігаючи всі повтори та вихідні позиції елементів кожного з джерел.')
        ],
        "⚡ Економія пам'яті:",
        "Поки не викликано .ToList() або .ToArray(), новий масив не виділяється в пам'яті."
    )
    add_code_panel(s76, 6.81, 1.55, 5.72, 4.25, [
        'IList<int> collection1 = new List<int>() { 1, 2, 3 };',
        'IList<int> collection2 = new List<int>() { 4, 5, 6 };',
        '',
        "// Об'єднання двох числових послідовностей",
        'var collection3 = collection1.Concat(collection2);',
        '',
        'foreach (int i in collection3)',
        '    Console.WriteLine(i);',
        '',
        "// Ланцюгове об'єднання з третьою колекцією:",
        'var collection4 = collection3.Concat(new List<int>() { 7, 8 });',
        'Console.WriteLine("Всього елементів: " + collection4.Count());',
        '',
        '-> Output:',
        '-> 1, 2, 3, 4, 5, 6',
        '-> Всього елементів: 8'
    ], 'Program.cs')
    add_bottom_banner(s76, "📌 Порядок слідування:",
                      "Гарантується строгий порядок: усі елементи першої колекції передують елементам другої колекції.")

    # ==========================================
    # SLIDE 77: SET OPERATORS ВСТУП ТА ОГЛЯД
    # ==========================================
    s77 = prs.slides.add_slide(blank_layout)
    add_bg(s77)
    add_header(s77, "14 · Оператори множин (Set Operators)", "Оператори роботи з множинами: Distinct, Except, Intersect, Union", 77)
    add_theory_card(s77, 0.8, 1.55, 5.72, 4.25,
        'Математичні операції над множинами в C#', ACCENT_INDIGO,
        [
            ('Чотири оператори множин:', "LINQ реалізує класичні теоретико-множинні операції: видалення дублікатів, різницю, перетин та об'єднання."),
            ('Усунення повторень:', 'Усі оператори над множинами повертають вибірки, де кожен унікальний елемент представлений лише один раз.'),
            ('Внутрішня реалізація:', 'Для швидкого пошуку дублікатів за O(1) LINQ використовує внутрішній хеш-набір Set<T>.'),
            ('Множинні перевантаження:', 'Кожен метод має перевантаження з IEqualityComparer<T> для порівняння складних сутностей за полями.')
        ],
        '⚡ Підтримка компараторів:',
        'Кожен оператор має перевантаження з параметром IEqualityComparer<T> для налаштування порівняння складних класів.'
    )
    add_theory_card(s77, 6.81, 1.55, 5.72, 4.25,
        "Призначення операторів множин", ACCENT_SKY,
        [
            ("Distinct:", "Повертає лише унікальні значення з однієї вихідної послідовності."),
            ("Except:", "Повертає різницю двох множин: елементи першої колекції, яких немає в другій."),
            ("Intersect:", "Повертає перетин двох множин: спільні елементи, які є в обох колекціях одночасно."),
            ("Union:", "Повертає математичне об'єднання двох колекцій без дублювання спільних елементів.")
        ],
        "📍 Deferred Execution:",
        "Усі оператори множин обчислюються відкладено і не матеріалізують результат до початку циклу або .ToList().",
        ACCENT_SKY
    )
    add_bottom_banner(s77, "💡 Теорія множин:",
                      "Оператори множин LINQ повністю повторюють операції над множинами в реляційних базах даних (DISTINCT, EXCEPT, INTERSECT, UNION).")

    # ==========================================
    # SLIDE 78: ТАБЛИЦЯ SET OPERATORS
    # ==========================================
    s78 = prs.slides.add_slide(blank_layout)
    add_bg(s78)
    add_header(s78, "14 · Оператори множин (Set Operators)", "Зведена порівняльна таблиця операторів над множинами", 78)
    headers_78 = ["Метод LINQ", "Опис / Призначення", "Що робить", "Приклад у C#", "Результат"]
    rows_78 = [
        ["Distinct", "Повертає унікальні значення з колекції.", "Видаляє всі дублікати з поточної послідовності.", "[1, 2, 2, 3].Distinct()", "[1, 2, 3]"],
        ["Except", "Повертає різницю між двома колекціями.", "Елементи першої колекції, яких немає в другій.", "[1, 2, 3].Except([2, 3, 4])", "[1]"],
        ["Intersect", "Повертає спільні елементи двох колекцій.", "Вибирає елементи, які є одночасно в обох списках.", "[1, 2, 3].Intersect([2, 3, 4])", "[2, 3]"],
        ["Union", "Об'єднує дві колекції без повторень.", "Всі елементи обох колекцій з усуненням дублікатів.", "[1, 2].Union([2, 3])", "[1, 2, 3]"]
    ]
    add_table_card(s78, 0.8, 1.55, 11.733, 2.15, headers_78, rows_78, [1.8, 2.5, 2.8, 2.7, 1.933], ACCENT_INDIGO)
    add_code_panel(s78, 0.8, 3.85, 11.733, 1.95, [
        '// Демонстрація операцій над множинами в C# (Set Operators):',
        'var colA = new int[] { 1, 2, 2, 3 };',
        'var colB = new int[] { 2, 3, 4 };',
        'var distinct  = colA.Distinct();           // -> [1, 2, 3] (дублікати числа 2 видалено)',
        'var except    = colA.Except(colB);         // -> [1] (елементи colA, яких немає в colB)',
        'var intersect = colA.Intersect(colB);      // -> [2, 3] (спільні елементи для обох колекцій)',
        "var union     = colA.Union(colB);          // -> [1, 2, 3, 4] (об'єднання без дублювання)",
        'Console.WriteLine($"Distinct:  {string.Join(", ", distinct)}");',
        'Console.WriteLine($"Except:    {string.Join(", ", except)}");',
        'Console.WriteLine($"Intersect: {string.Join(", ", intersect)}");',
        'Console.WriteLine($"Union:     {string.Join(", ", union)}");'
    ], 'SetOperatorsDemo.cs')
    add_bottom_banner(s78, "💡 Резюме поведінки:",
                      "Усі 4 методи повертають IEnumerable<T> та гарантують абсолютну унікальність кожного значення у вихідній послідовності.")

    # ==========================================
    # SLIDE 79: ВІЗУАЛЬНА СХЕМА РОБОТИ SET OPERATORS
    # ==========================================
    s79 = prs.slides.add_slide(blank_layout)
    add_bg(s79)
    add_header(s79, "14 · Оператори множин (Set Operators)", "Візуалізація операцій на прикладі тестових множин", 79)
    add_dual_comparison_card(s79, 0.8, 1.55, 5.72, 4.25,
        "Distinct та Except", ACCENT_INDIGO,
        [
            ("Вхідна колекція 1:", "Col1 = { \"One\", \"Two\", \"Three\", \"Two\" }"),
            ("Distinct(Col1):", "Видаляє другий \"Two\" -> { \"One\", \"Two\", \"Three\" }"),
            ("Except(Col1, Col2):", "Де Col2 = { \"Two\", \"Four\" } -> повертає { \"One\", \"Three\" } (елементи з Col1, яких нема в Col2)")
        ],
        [
            "// 1. Distinct - усунення повторів:",
            "var dist = col1.Distinct();",
            "// -> One, Two, Three",
            "",
            "// 2. Except - різниця множин:",
            "var exc = col1.Except(col2);",
            "// -> One, Three"
        ]
    )
    add_dual_comparison_card(s79, 6.81, 1.55, 5.72, 4.25,
        "Intersect та Union", ACCENT_SKY,
        [
            ("Вхідна колекція 2:", "Col2 = { \"Two\", \"Four\", \"Five\" }"),
            ("Intersect(Col1, Col2):", "Спільні елементи для обох списків -> { \"Two\" }"),
            ("Union(Col1, Col2):", "Повне об'єднання всіх унікальних -> { \"One\", \"Two\", \"Three\", \"Four\", \"Five\" }")
        ],
        [
            "// 3. Intersect - перетин множин:",
            "var inter = col1.Intersect(col2);",
            "// -> Two",
            "",
            "// 4. Union - об'єднання без дублікатів:",
            "var un = col1.Union(col2);",
            "// -> One, Two, Three, Four, Five"
        ]
    )
    add_bottom_banner(s79, "📌 Комутативність:",
                      "Intersect та Union є комутативними за складом (A та B дають однакові множини). Except НЕ є комутативним: A \\ B != B \\ A.")

    # ==========================================
    # SLIDE 80: DISTINCT НА ПРОСТИХ ТИПАХ
    # ==========================================
    s80 = prs.slides.add_slide(blank_layout)
    add_bg(s80)
    add_header(s80, "14 · Оператори множин (Set Operators)", "Оператор Distinct: Очищення рядків та чисел від дублікатів", 80)
    add_theory_card(s80, 0.8, 1.55, 5.72, 4.25,
        'Фільтрація унікальних значень', ACCENT_INDIGO,
        [
            ('Призначення:', 'Distinct повертає кожне неповторюване значення з колекції рівно один раз.'),
            ('Збереження порядку:', 'Порядок першої появи кожного елемента гарантовано зберігається у вихідному результаті.'),
            ('Чутливість до регістру:', "Для рядків за замовчуванням 'Three' і 'three' є різними значеннями. Для нечутливого порівняння передають StringComparer.OrdinalIgnoreCase."),
            ('Потокова фільтрація:', 'Distinct віддає унікальні елементи на льоту через yield return, як тільки зустрічає нове неповторюване значення.')
        ],
        '⚡ Результат для списків:',
        'Рядки: One, Two, Three (повтори відкинуто). Числа: 1, 2, 3, 4, 5 (повтори чисел 2, 4, 3 відкинуто).'
    )
    add_code_panel(s80, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList = new List<string>() {',
        '    "One", "Two", "Three", "Two", "Three"',
        '};',
        'IList<int> intList = new List<int>() { 1, 2, 3, 2, 4, 4, 3, 5 };',
        '',
        'var distinctList1 = strList.Distinct();',
        'foreach (var str in distinctList1)',
        '    Console.WriteLine(str);',
        '',
        'var distinctList2 = intList.Distinct();',
        'foreach (var i in distinctList2)',
        '    Console.WriteLine(i);',
        '',
        '-> Output рядків: One, Two, Three',
        '-> Output чисел: 1, 2, 3, 4, 5'
    ], 'Program.cs')
    add_bottom_banner(s80, "💡 Складність операції:",
                      "Distinct працює з часовою складністю O(N) завдяки використанню внутрішньої хеш-таблиці.")

    # ==========================================
    # SLIDE 81: DISTINCT ДЛЯ ОБ'ЄКТІВ STUDENT
    # ==========================================
    s81 = prs.slides.add_slide(blank_layout)
    add_bg(s81)
    add_header(s81, "14 · Оператори множин (Set Operators)", "Оператор Distinct для об'єктів Student з StudentComparer", 81)
    add_theory_card(s81, 0.8, 1.55, 5.72, 4.25,
        "Усунення дублікатів складних об'єктів", ACCENT_INDIGO,
        [
            ('Проблема посилань:', "У списку є 4 студенти з ім'ям 'Bill', ID = 3, Age = 25. Без компаратора Distinct поверне всі 4 об'єкти, бо це різні адреси в купі!"),
            ('Передача StudentComparer:', 'Передаючи екземпляр IEqualityComparer<Student>, ми змушуємо Distinct порівнювати StudentID та StudentName.'),
            ('Сучасна альтернатива:', 'У C# 9+ використання record struct або record class забезпечує автопорівняння за значеннями без окремих компараторів.'),
            ("Хешування об'єктів:", 'Distinct спочатку викликає GetHashCode, а лише при колізії хешів — дорожчий метод Equals.')
        ],
        '⚡ Результат фільтрації:',
        'Повертаються лише 4 унікальні студенти: John, Steve, Bill (один раз) та Ron.'
    )
    add_code_panel(s81, 6.81, 1.55, 5.72, 4.25, [
        'IList<Student> studentList = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        '',
        'var distinctStudents = studentList.Distinct(',
        '    new StudentComparer());',
        '',
        'foreach (Student std in distinctStudents)',
        '    Console.WriteLine(std.StudentName);',
        '',
        '-> Output: John, Steve, Bill, Ron'
    ], 'StudentComparer.cs')
    add_bottom_banner(s81, "📌 Зверніть увагу:",
                      "У вихідний результат потрапляє перший зустрінутий екземпляр з групи еквівалентних об'єктів.")

    # ==========================================
    # SLIDE 82: EXCEPT ДЛЯ РЯДКОВИХ КОЛЕКЦІЙ
    # ==========================================
    s82 = prs.slides.add_slide(blank_layout)
    add_bg(s82)
    add_header(s82, "14 · Оператори множин (Set Operators)", "Оператор Except: Різниця між двома колекціями рядків", 82)
    add_theory_card(s82, 0.8, 1.55, 5.72, 4.25,
        'Обчислення різниці послідовностей (A \\ B)', ACCENT_INDIGO,
        [
            ('Принцип Except:', 'Повертає елементи першої колекції, які відсутні в другій колекції (параметрі).'),
            ('Видалення збігів:', "Елементи 'Four' та 'Five' є в обох списках, тому вони повністю вилучаються з результату."),
            ('Елементи другої колекції:', "Елементи 'Six', 'Seven', 'Eight' із strList2 взагалі не потрапляють у результат, бо їх не було в strList1."),
            ('Одноразове сканування:', 'Друга колекція зчитується в хеш-сет Set<T>, після чого перший список фільтрується за лінійний час O(N + M).')
        ],
        '⚡ Результат вибірки:',
        'Повертаються тільки: One, Two, Three. Усі спільні та сторонні значення відсіяні.'
    )
    add_code_panel(s82, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList1 = new List<string>() {',
        '    "One", "Two", "Three", "Four", "Five"',
        '};',
        '',
        'IList<string> strList2 = new List<string>() {',
        '    "Four", "Five", "Six", "Seven", "Eight"',
        '};',
        '',
        '// Елементи strList1, яких немає в strList2',
        'var result = strList1.Except(strList2);',
        '',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        '-> Output:',
        '-> One',
        '-> Two',
        '-> Three'
    ], 'Program.cs')
    add_bottom_banner(s82, "💡 Некомутативність:",
                      "strList2.Except(strList1) повернув би зовсім інший результат: Six, Seven, Eight.")

    # ==========================================
    # SLIDE 83: EXCEPT ДЛЯ ОБ'ЄКТІВ STUDENT
    # ==========================================
    s83 = prs.slides.add_slide(blank_layout)
    add_bg(s83)
    add_header(s83, "14 · Оператори множин (Set Operators)", "Оператор Except для об'єктів Student з StudentComparer", 83)
    add_theory_card(s83, 0.8, 1.55, 5.72, 4.25,
        "Вилучення підмножини бізнес-об'єктів", ACCENT_INDIGO,
        [
            ('Вихідний список 1:', 'Містить студентів: John (1), Steve (2), Bill (3), Ron (5).'),
            ('Список вилучення 2:', 'Містить студентів: Bill (3), Ron (5).'),
            ('Порівняння компаратором:', 'Завдяки StudentComparer метод Except знаходить відповідність за ID та іменем і коректно віднімає студентів.'),
            ('Ідентифікація сутностей:', "StudentComparer зіставляє ID та ім'я, забезпечуючи коректне вилучення студентів із бази.")
        ],
        '⚡ Підсумок різниці:',
        'У результуючій колекції залишаються тільки John та Steve.'
    )
    add_code_panel(s83, 6.81, 1.55, 5.72, 4.25, [
        'IList<Student> studentList1 = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        'IList<Student> studentList2 = new List<Student>() {',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        '',
        'var resultedCol = studentList1.Except(',
        '    studentList2, new StudentComparer());',
        '',
        'foreach (Student std in resultedCol)',
        '    Console.WriteLine(std.StudentName);',
        '',
        '-> Output: John, Steve'
    ], 'StudentComparer.cs')
    add_bottom_banner(s83, "📌 Практичне застосування:",
                      "Швидкий пошук неактивних користувачів, боржників або студентів, які ще не записані на обрану дисципліну.")

    # ==========================================
    # SLIDE 84: INTERSECT ДЛЯ РЯДКОВИХ КОЛЕКЦІЙ
    # ==========================================
    s84 = prs.slides.add_slide(blank_layout)
    add_bg(s84)
    add_header(s84, "14 · Оператори множин (Set Operators)", "Оператор Intersect: Перетин двох списків рядків", 84)
    add_theory_card(s84, 0.8, 1.55, 5.72, 4.25,
        'Пошук спільних елементів (A ∩ B)', ACCENT_INDIGO,
        [
            ('Принцип Intersect:', 'Метод повертає нову колекцію, що містить тільки ті елементи, які присутні в обох джерелах одночасно.'),
            ('Унікальність результату:', 'У вихідній послідовності гарантовано відсутні дублікати, навіть якщо в джерелах спільний елемент повторювався.'),
            ('Порядок елементів:', 'Порядок елементів визначається порядком їх появи в першій колекції (strList1).'),
            ('Хеш-перетин:', 'LINQ створює хеш-набір з другої колекції і повертає елементи першої, що знайдені в ньому, видаляючи їх для уникнення дублікатів.')
        ],
        '⚡ Результат перетину:',
        "Спільними для двох списків є лише 'Four' та 'Five'."
    )
    add_code_panel(s84, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList1 = new List<string>() {',
        '    "One", "Two", "Three", "Four", "Five"',
        '};',
        '',
        'IList<string> strList2 = new List<string>() {',
        '    "Four", "Five", "Six", "Seven", "Eight"',
        '};',
        '',
        '// Пошук спільних елементів',
        'var result = strList1.Intersect(strList2);',
        '',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        '-> Output:',
        '-> Four',
        '-> Five'
    ], 'Program.cs')
    add_bottom_banner(s84, "💡 Бізнес-застосування:",
                      "Перевірка перетину прав доступу користувача, знаходження спільних спільників або спільних тегів товарів.")

    # ==========================================
    # SLIDE 85: INTERSECT ДЛЯ ОБ'ЄКТІВ STUDENT
    # ==========================================
    s85 = prs.slides.add_slide(blank_layout)
    add_bg(s85)
    add_header(s85, "14 · Оператори множин (Set Operators)", "Оператор Intersect для об'єктів Student з StudentComparer", 85)
    add_theory_card(s85, 0.8, 1.55, 5.72, 4.25,
        'Знаходження спільних студентів у двох групах', ACCENT_INDIGO,
        [
            ('Колекція 1:', 'John (1), Steve (2), Bill (3), Ron (5).'),
            ('Колекція 2:', 'Bill (3), Ron (5).'),
            ('Перетин списків:', 'Intersect порівнює сутності через StudentComparer і повертає спільних студентів.'),
            ('Спільний реєстр:', 'Знаходить тих самих студентів у двох різних академічних групах або списках курсів.')
        ],
        '⚡ Результат вибірки:',
        "Повертаються об'єкти Bill та Ron, які присутні в обох колекціях одночасно."
    )
    add_code_panel(s85, 6.81, 1.55, 5.72, 4.25, [
        'IList<Student> studentList1 = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        'IList<Student> studentList2 = new List<Student>() {',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        '',
        'var resultedCol = studentList1.Intersect(',
        '    studentList2, new StudentComparer());',
        '',
        'foreach (Student std in resultedCol)',
        '    Console.WriteLine(std.StudentName);',
        '',
        '-> Output: Bill, Ron'
    ], 'StudentComparer.cs')
    add_bottom_banner(s85, "📌 Джерело екземплярів:",
                      "Об'єкти у вихідній послідовності повертаються з першої колекції (studentList1).")

    # ==========================================
    # SLIDE 86: UNION ДЛЯ РЯДКОВИХ КОЛЕКЦІЙ
    # ==========================================
    s86 = prs.slides.add_slide(blank_layout)
    add_bg(s86)
    add_header(s86, "14 · Оператори множин (Set Operators)", "Оператор Union: Математичне об'єднання списків рядків", 86)
    add_theory_card(s86, 0.8, 1.55, 5.72, 4.25,
        "Об'єднання множин без дублювання (A ∪ B)", ACCENT_INDIGO,
        [
            ('Принцип Union:', "Об'єднує дві послідовності та повертає всі унікальні елементи обох списків."),
            ('Видалення дублікатів:', "Значення 'Two' є в обох списках, але в результаті з'являється лише один раз."),
            ('Чутливість до регістру:', "'three' та 'THREE' розглядаються як різні рядки і обидва потрапляють у вибірку."),
            ('Ліниве злиття:', 'Union послідовно читає обидва джерела, видаючи кожен новий неповторюваний елемент через yield return.')
        ],
        '⚡ Результат вибірки:',
        'One, Two, three, THREE, Four, Five. Забезпечено повну унікальність значень.'
    )
    add_code_panel(s86, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList1 = new List<string>() {',
        '    "One", "Two", "three"',
        '};',
        '',
        'IList<string> strList2 = new List<string>() {',
        '    "Two", "THREE", "Four", "Five"',
        '};',
        '',
        "// Об'єднання двох множин з усуненням дублікатів",
        'var result = strList1.Union(strList2);',
        '',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        '-> Output:',
        '-> One, Two, three, THREE, Four, Five'
    ], 'Program.cs')
    add_bottom_banner(s86, "💡 Порівняння з Concat:",
                      "Якби замість Union було викликано Concat, 'Two' повторилося б двічі.")

    # ==========================================
    # SLIDE 87: UNION ДЛЯ ОБ'ЄКТІВ STUDENT
    # ==========================================
    s87 = prs.slides.add_slide(blank_layout)
    add_bg(s87)
    add_header(s87, "14 · Оператори множин (Set Operators)", "Оператор Union для об'єктів Student з StudentComparer", 87)
    add_theory_card(s87, 0.8, 1.55, 5.72, 4.25,
        "Об'єднання реєстрів студентів без дублювання", ACCENT_INDIGO,
        [
            ('Колекція 1:', 'John (1), Steve (2), Bill (3), Ron (5).'),
            ('Колекція 2:', 'Bill (3), Ron (5).'),
            ('Усунення повторів:', 'Оскільки Bill та Ron вже є в studentList1, метод Union завдяки StudentComparer не додає їх повторно.'),
            ('Консолідація баз даних:', 'Ідеально підходить для дедуплікації даних при синхронізації кількох локальних реєстрів в один.')
        ],
        "⚡ Результат об'єднання:",
        'John, Steve, Bill, Ron — єдиний консолідований список унікальних студентів.'
    )
    add_code_panel(s87, 6.81, 1.55, 5.72, 4.25, [
        'IList<Student> studentList1 = new List<Student>() {',
        '    new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '    new Student() { StudentID = 2, StudentName = "Steve", Age = 15 },',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        'IList<Student> studentList2 = new List<Student>() {',
        '    new Student() { StudentID = 3, StudentName = "Bill", Age = 25 },',
        '    new Student() { StudentID = 5, StudentName = "Ron", Age = 19 }',
        '};',
        '',
        'var resultedCol = studentList1.Union(',
        '    studentList2, new StudentComparer());',
        '',
        'foreach (Student std in resultedCol)',
        '    Console.WriteLine(std.StudentName);',
        '',
        '-> Output: John, Steve, Bill, Ron'
    ], 'StudentComparer.cs')
    add_bottom_banner(s87, "📌 Підсумок блоку Set Operators:",
                      "Усі 4 оператори множин (Distinct, Except, Intersect, Union) вимагають компаратора при роботі з користувацькими класами.")

    # ==========================================
    # SLIDE 88: ТАБЛИЦЯ PARTITIONING OPERATORS
    # ==========================================
    s88 = prs.slides.add_slide(blank_layout)
    add_bg(s88)
    add_header(s88, "15 · Секціонування (Partitioning)", "Оператори розбиття послідовностей: Skip, SkipWhile, Take, TakeWhile", 88)
    headers_88 = ["Метод LINQ", "Опис / Призначення", "Що робить", "Приклад у C#", "Результат"]
    rows_88 = [
        ["Skip", "Пропускає задану кількість елементів з початку послідовності.", "Пропускає перші N елементів і повертає всі наступні.", "[1, 2, 3, 4].Skip(2)", "[3, 4]"],
        ["SkipWhile", "Пропускає елементи з початку, поки умова true.", "Як тільки умова стає false — припиняє пропуск і повертає решту послідовності.", "[1, 2, 3, 4].SkipWhile(x => x < 3)", "[3, 4]"],
        ["Take", "Бере задану кількість елементів з початку послідовності.", "Повертає тільки перші N елементів і завершує вибірку.", "[1, 2, 3, 4].Take(2)", "[1, 2]"],
        ["TakeWhile", "Бере елементи з початку, поки умова true.", "Як тільки умова стає false — негайно зупиняє подальшу вибірку.", "[1, 2, 3, 4].TakeWhile(x => x < 3)", "[1, 2]"]
    ]
    add_table_card(s88, 0.8, 1.55, 11.733, 2.15, headers_88, rows_88, [1.7, 2.6, 2.7, 2.8, 1.933], ACCENT_INDIGO)
    add_code_panel(s88, 0.8, 3.85, 11.733, 1.95, [
        '// Канонічний патерн сторінкової пагінації (Pagination Pattern):',
        'int pageNumber = 2, pageSize = 3;',
        'var items = Enumerable.Range(1, 10); // послідовність чисел 1..10',
        'var page = items.Skip((pageNumber - 1) * pageSize).Take(pageSize);',
        'Console.WriteLine($"Сторінка {pageNumber}: {string.Join(", ", page)}"); // 4, 5, 6',
        '',
        '// Умовне позиційне секціонування (TakeWhile / SkipWhile):',
        'var whileTake = items.TakeWhile(x => x < 4); // [1, 2, 3] (до першого false)',
        'var whileSkip = items.SkipWhile(x => x < 7); // [7, 8, 9, 10] (пропуск до першого false)',
        'Console.WriteLine($"TakeWhile: {string.Join(", ", whileTake)}");',
        'Console.WriteLine($"SkipWhile: {string.Join(", ", whileSkip)}");'
    ], 'PaginationPattern.cs')
    add_bottom_banner(s88, "💡 Патерн Pagination:",
                      "Спільне використання: source.Skip((pageNumber - 1) * pageSize).Take(pageSize) реалізує сторінкову розбивку в UI та API.")

    # ==========================================
    # SLIDE 89: SKIP
    # ==========================================
    s89 = prs.slides.add_slide(blank_layout)
    add_bg(s89)
    add_header(s89, "15 · Секціонування (Partitioning)", "Оператор Skip: Пропуск фіксованої кількості початкових елементів", 89)
    add_theory_card(s89, 0.8, 1.55, 5.72, 4.25,
        'Пропуск перших N елементів', ACCENT_INDIGO,
        [
            ('Принцип роботи:', 'Skip(count) ігнорує перші count елементів у послідовності та повертає всі інші.'),
            ('Якщо count <= 0:', 'Метод повертає всю вихідну послідовність без змін.'),
            ('Якщо count >= розміру:', 'Повертається порожня колекція (виняток IndexOutOfRange НЕ викидається).'),
            ('Трансляція в SQL:', 'У базах даних транслюється у вираз OFFSET n ROWS.')
        ],
        '⚡ Результат для списку:',
        "Після пропуску перших 2 рядків ('One', 'Two') повертаються: Three, Four, Five."
    )
    add_code_panel(s89, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList = new List<string>() {',
        '    "One", "Two", "Three", "Four", "Five"',
        '};',
        '',
        '// Пропускаємо перші 2 елементи колекції',
        'var newList = strList.Skip(2);',
        '',
        'Console.WriteLine("--- Результат Skip(2) ---");',
        'foreach (var str in newList)',
        '    Console.WriteLine(str);',
        '',
        '// Пропуск більше ніж є елементів (порожній результат):',
        'var empty = strList.Skip(10);',
        'Console.WriteLine("Кількість у empty: " + empty.Count());',
        '',
        '-> Output:',
        '-> Three',
        '-> Four',
        '-> Five',
        '-> Кількість у empty: 0'
    ], 'Program.cs')
    add_bottom_banner(s89, "📌 Ефективність:",
                      "Skip виконується відкладено і не матеріалізує пропущені елементи в оперативну пам'ять.")

    # ==========================================
    # SLIDE 90: SKIPWHILE БАЗОВИЙ
    # ==========================================
    s90 = prs.slides.add_slide(blank_layout)
    add_bg(s90)
    add_header(s90, "15 · Секціонування (Partitioning)", "Оператор SkipWhile: Пропуск за предикатною умовою", 90)
    add_theory_card(s90, 0.8, 1.55, 5.72, 4.25,
        "Пропуск елементів до першого false", ACCENT_INDIGO,
        [
            ("Предикатна перевірка:", "SkipWhile обчислює умову s.Length < 4 для кожного елемента послідовно."),
            ("Елементи 1 і 2:", "'One' (довжина 3 < 4: true -> пропуск), 'Two' (довжина 3 < 4: true -> пропуск)."),
            ("Точка зупинки:", "На елементі 'Three' довжина дорівнює 5 (5 < 4: false). Пропуск негайно зупиняється!"),
            ("Решта колекції:", "Усі наступні елементи ('Four', 'Five', 'Six') повертаються без повторної перевірки умови.")
        ],
        "⚡ Результат вибірки:",
        "Повертаються: Three, Four, Five, Six."
    )
    add_code_panel(s90, 6.81, 1.55, 5.72, 4.25, [
        "IList<string> strList = new List<string>() {",
        "    \"One\", \"Two\", \"Three\", \"Four\", \"Five\", \"Six\"",
        "};",
        "",
        "// Пропускаємо, поки довжина слова менша за 4",
        "var resultList = strList.SkipWhile(s => s.Length < 4);",
        "",
        "foreach (string str in resultList)",
        "    Console.WriteLine(str);",
        "",
        "-> Output:",
        "-> Three",
        "-> Four",
        "-> Five",
        "-> Six"
    ], "Program.cs")
    add_bottom_banner(s90, "💡 Різниця з Where:",
                      "Where відфільтрував би й слово 'Six' (довжина 3), а SkipWhile повертає 'Six', бо після першого false перевірка припинилася!")

    # ==========================================
    # SLIDE 91: SKIPWHILE КОЛИ ПЕРШИЙ FALSE
    # ==========================================
    s91 = prs.slides.add_slide(blank_layout)
    add_bg(s91)
    add_header(s91, "15 · Секціонування (Partitioning)", "SkipWhile: Поведінка, коли найперший елемент не задовольняє умову", 91)
    add_theory_card(s91, 0.8, 1.55, 5.72, 4.25,
        "Миттєве зупинення пропуску на позиції 0", ACCENT_INDIGO,
        [
            ("Аналіз першого елемента:", "Перший елемент списку — 'Three' (його довжина 5)."),
            ("Умова s.Length < 4:", "5 < 4 одразу дає значення false!"),
            ("Результат пропуску:", "Оскільки перший же елемент не задовольнив умову, пропускається рівно 0 елементів."),
            ("Повернення даних:", "Повертається абсолютно вся послідовність у повному обсязі, починаючи з найпершого елемента.")
        ],
        "⚡ Результат вибірки:",
        "Three, One, Two, Four, Five, Six — жоден елемент не пропущено!"
    )
    add_code_panel(s91, 6.81, 1.55, 5.72, 4.25, [
        "IList<string> strList = new List<string>() {",
        "    \"Three\", \"One\", \"Two\", \"Four\", \"Five\", \"Six\"",
        "};",
        "",
        "// Перший елемент 'Three' має довжину 5 >= 4",
        "// Умова s.Length < 4 дає false одразу:",
        "var resultList = strList.SkipWhile(s => s.Length < 4);",
        "",
        "foreach (string str in resultList)",
        "    Console.WriteLine(str);",
        "",
        "-> Output:",
        "-> Three",
        "-> One",
        "-> Two",
        "-> Four",
        "-> Five",
        "-> Six"
    ], "Program.cs")
    add_bottom_banner(s91, "📌 Головний висновок:",
                      "SkipWhile не шукає перше співпадіння вглибині колекції — він перевіряє елементи виключно від самого початку.")

    # ==========================================
    # SLIDE 92: SKIPWHILE З ІНДЕКСОМ
    # ==========================================
    s92 = prs.slides.add_slide(blank_layout)
    add_bg(s92)
    add_header(s92, "15 · Секціонування (Partitioning)", "Оператор SkipWhile з використанням індексу елемента", 92)
    add_theory_card(s92, 0.8, 1.55, 5.72, 4.25,
        'Лямбда з парою (елемент, індекс)', ACCENT_INDIGO,
        [
            ('Перевантаження з індексом:', 'SkipWhile((s, i) => s.Length > i) передає 0-based індекс кожного елемента.'),
            ('Трасування кроків:', "[0] 'One' (3 > 0: true), [1] 'Two' (3 > 1: true), [2] 'Three' (5 > 2: true), [3] 'Four' (4 > 3: true)."),
            ('Критична точка [4]:', "Елемент 'Five' (довжина 4): умова 4 > 4 стає false!"),
            ('Зупинка пропуску:', "Починаючи з індексу 4 ('Five') і до кінця списку, всі елементи повертаються.")
        ],
        '⚡ Результат вибірки:',
        'Повертаються лише останні елементи: Five та Six.'
    )
    add_code_panel(s92, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList = new List<string>() {',
        '    "One", "Two", "Three", "Four", "Five", "Six"',
        '};',
        '',
        '// Пропускаємо, поки довжина рядка більша за його індекс:',
        "// [0]'One'(3>0:T), [1]'Two'(3>1:T), [2]'Three'(5>2:T), [3]'Four'(4>3:T)",
        "// [4]'Five'(4>4: false -> зупинка пропуску!)",
        'var result = strList.SkipWhile((s, i) => s.Length > i);',
        '',
        'Console.WriteLine("--- Елементи після зупинки пропуску ---");',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        '-> Output:',
        '-> Five',
        '-> Six'
    ], 'Program.cs')
    add_bottom_banner(s92, "💡 Застосування індексів:",
                      "Індексовані предикати дозволяють реалізовувати алгоритми нелінійного ковзного вікна або складного позиційного розбору.")

    # ==========================================
    # SLIDE 93: TAKE
    # ==========================================
    s93 = prs.slides.add_slide(blank_layout)
    add_bg(s93)
    add_header(s93, "15 · Секціонування (Partitioning)", "Оператор Take: Отримання перших N елементів послідовності", 93)
    add_theory_card(s93, 0.8, 1.55, 5.72, 4.25,
        'Вибірка фіксованої кількості елементів', ACCENT_INDIGO,
        [
            ('Принцип роботи:', 'Take(count) повертає рівно count перших елементів від початку і негайно завершує перебір.'),
            ('Раннє завершення:', 'Метод не сканує всю колекцію до кінця, що критично важливо для швидкодії.'),
            ('Якщо count <= 0:', 'Повертається порожня послідовність.'),
            ('Трансляція в SQL:', 'У базах даних транслятор LINQ генерує TOP (count) або FETCH NEXT count ROWS ONLY.')
        ],
        '⚡ Результат для списку:',
        'Take(2) повертає перші два елементи: One та Two.'
    )
    add_code_panel(s93, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList = new List<string>() {',
        '    "One", "Two", "Three", "Four", "Five"',
        '};',
        '',
        '// Отримуємо рівно 2 перших елементи послідовності',
        'var result = strList.Take(2);',
        '',
        'Console.WriteLine("--- Перші 2 елементи ---");',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        '// Запит більшої кількості ніж є в наявності:',
        'var all = strList.Take(10);',
        'Console.WriteLine("Всього взято: " + all.Count());',
        '',
        '-> Output:',
        '-> One',
        '-> Two',
        '-> Всього взято: 5'
    ], 'Program.cs')
    add_bottom_banner(s93, "📌 Ранній вихід:",
                      "Take негайно зупиняє ітератор MoveNext() після отримання потрібної кількості елементів.")

    # ==========================================
    # SLIDE 94: TAKEWHILE БАЗОВИЙ
    # ==========================================
    s94 = prs.slides.add_slide(blank_layout)
    add_bg(s94)
    add_header(s94, "15 · Секціонування (Partitioning)", "Оператор TakeWhile: Вибірка елементів до першого порушення умови", 94)
    add_theory_card(s94, 0.8, 1.55, 5.72, 4.25,
        'Отримання елементів, поки умова true', ACCENT_INDIGO,
        [
            ('Принцип TakeWhile:', 'Бере елементи від початку списку, поки предикатна умова s.Length > 4 повертає true.'),
            ("Крок 1 'Three':", 'Довжина 5 > 4: true -> елемент береться.'),
            ("Крок 2 'Four':", 'Довжина 4 > 4: false -> вибірка негайно припиняється!'),
            ('Ігнорування решти:', "Навіть якщо далі є елемент 'Hundred' з довжиною 7 (> 4), він ніколи не буде оброблений.")
        ],
        '⚡ Результат вибірки:',
        'Повертається лише один елемент: Three.'
    )
    add_code_panel(s94, 6.81, 1.55, 5.72, 4.25, [
        'IList<string> strList = new List<string>() {',
        '    "Three", "Four", "Five", "Hundred"',
        '};',
        '',
        '// Беремо елементи, поки їх довжина більша за 4:',
        "// [0]'Three'(len=5 > 4: true) -> взято",
        "// [1]'Four' (len=4 > 4: false) -> СТОП!",
        'var result = strList.TakeWhile(s => s.Length > 4);',
        '',
        'Console.WriteLine("--- Результат TakeWhile ---");',
        'foreach (string str in result)',
        '    Console.WriteLine(str);',
        '',
        "// 'Hundred' (len=7) ніколи не буде прочитано!",
        '-> Output:',
        '-> Three'
    ], 'Program.cs')
    add_bottom_banner(s94, "💡 Чому 'Hundred' не потрапив у результат:",
                      "TakeWhile повністю зупиняє ітерування, як тільки умова порушилася на слові 'Four'. Для повної фільтрації використовуйте Where.")

    # ==========================================
    # SLIDE 95: TAKEWHILE З ІНДЕКСОМ
    # ==========================================
    s95 = prs.slides.add_slide(blank_layout)
    add_bg(s95)
    add_header(s95, "15 · Секціонування (Partitioning)", "Оператор TakeWhile з використанням індексу елемента", 95)
    add_theory_card(s95, 0.8, 1.55, 5.72, 4.25,
        "Індексована зупинка вибірки", ACCENT_INDIGO,
        [
            ("Умова s.Length > i:", "Порівнює довжину поточного рядка з його порядковим індексом."),
            ("[0] 'One' (3 > 0):", "true -> беремо"),
            ("[1] 'Two' (3 > 1):", "true -> беремо"),
            ("[2] 'Three' (5 > 2):", "true -> беремо"),
            ("[3] 'Four' (4 > 3):", "true -> беремо"),
            ("[4] 'Five' (4 > 4):", "false -> фінальна зупинка!")
        ],
        "⚡ Результат вибірки:",
        "Отримано перші 4 елементи: One, Two, Three, Four."
    )
    add_code_panel(s95, 6.81, 1.55, 5.72, 4.25, [
        "IList<string> strList = new List<string>() {",
        "    \"One\", \"Two\", \"Three\", \"Four\", \"Five\", \"Six\"",
        "};",
        "",
        "// Беремо, поки довжина рядка більша за його індекс",
        "var resultList = strList.TakeWhile((s, i) => s.Length > i);",
        "",
        "foreach (string str in resultList)",
        "    Console.WriteLine(str);",
        "",
        "-> Output:",
        "-> One",
        "-> Two",
        "-> Three",
        "-> Four"
    ], "Program.cs")
    add_bottom_banner(s95, "📌 Резюме блоку Partitioning:",
                      "Skip, SkipWhile, Take та TakeWhile забезпечують повний контроль над порційним позиційним зчитуванням послідовностей.")

    # ==========================================
    # SLIDE 96: ТАБЛИЦЯ CONVERSION OPERATORS
    # ==========================================
    s96 = prs.slides.add_slide(blank_layout)
    add_bg(s96)
    add_header(s96, "16 · Перетворення (Conversion)", "Оператори приведення та матеріалізації типів (Conversion Operators)", 96)
    headers_96 = ["Метод LINQ", "Категорія", "Опис / Призначення", "Що повертає"]
    rows_96 = [
        ["AsEnumerable", "As-оператор", "Приховує кастомну реалізацію провайдера, повертаючи IEnumerable<T>.", "IEnumerable<T>"],
        ["AsQueryable", "As-оператор", "Перетворює IEnumerable на IQueryable для емуляції провайдера запитів.", "IQueryable<T>"],
        ["Cast<T>", "Cast-оператор", "Приводить елементи нетипізованої колекції (ArrayList) до узагальненого типу T.", "IEnumerable<T>"],
        ["OfType<T>", "Cast-оператор", "Безпечно фільтрує колекцію, вибираючи лише елементи вказаного типу T.", "IEnumerable<T>"],
        ["ToArray", "To-оператор", "Негайно виконує запит і копіює всі елементи в одновимірний масив T[].", "T[]"],
        ["ToList", "To-оператор", "Негайно виконує запит і створює новий екземпляр List<T> у пам'яті.", "List<T>"],
        ["ToDictionary", "To-оператор", "Матеріалізує елементи в словник Dictionary<TKey, TValue> за ключем.", "Dictionary<K, V>"],
        ["ToLookup", "To-оператор", "Групує елементи в незмінну 1:N структуру Lookup<TKey, TElement>.", "ILookup<K, E>"]
    ]
    add_table_card(s96, 0.8, 1.55, 11.733, 4.25, headers_96, rows_96, [2.0, 1.8, 5.733, 2.2], ACCENT_INDIGO)
    add_bottom_banner(s96, "💡 Три групи перетворень:",
                      "As-оператори змінюють інтерфейс доступу; Cast-оператори валідують типи; To-оператори примусово матеріалізують дані (Immediate Execution).")

    # ==========================================
    # SLIDE 97: ASENUMERABLE ТА ASQUERYABLE
    # ==========================================
    s97 = prs.slides.add_slide(blank_layout)
    add_bg(s97)
    add_header(s97, "16 · Перетворення (Conversion)", "Оператори AsEnumerable та AsQueryable: Статичні та динамічні типи", 97)
    add_theory_card(s97, 0.8, 1.55, 5.72, 4.25,
        'Зміна інтерфейсу представлення колекції', ACCENT_INDIGO,
        [
            ('AsEnumerable():', "Переводить статичний тип колекції в IEnumerable<T>. Критично у LINQ to Entities: дозволяє примусово переключити подальшу частину запиту на виконання в пам'яті C# замість SQL сервера."),
            ('AsQueryable():', 'Огортає звичайний IEnumerable у EnumerableQuery<T>, дозволяючи імітувати роботу IQueryable провайдера (наприклад, у Unit-тестах).'),
            ('Типи компіляції та виконання:', 'Демонстрація різниці між typeof(T).Name та obj.GetType().Name під час виклику утиліти.'),
            ('Клієнтська оцінка (Client Evaluation):', "AsEnumerable перемикає виконання запиту в пам'ять процесу для виклику методів C#, які не підтримуються SQL.")
        ],
        '⚡ Важливий нюанс:',
        "AsEnumerable НЕ створює копії масиву — у рантаймі об'єкт залишається тим самим вихідним Student[]."
    )
    add_code_panel(s97, 6.81, 1.55, 5.72, 4.25, [
        'class Program {',
        '    static void ReportTypeProperties<T>(T obj) {',
        '        Console.WriteLine("Compile-time: " + typeof(T).Name);',
        '        Console.WriteLine("Actual runtime: " + obj.GetType().Name);',
        '    }',
        '    static void Main(string[] args) {',
        '        Student[] studentArray = {',
        '            new Student() { StudentID = 1, StudentName = "John", Age = 18 },',
        '            new Student() { StudentID = 2, StudentName = "Steve", Age = 21 }',
        '        };',
        '        ReportTypeProperties(studentArray);',
        '        ReportTypeProperties(studentArray.AsEnumerable());',
        '        ReportTypeProperties(studentArray.AsQueryable());',
        '    }',
        '}',
        '-> Compile-time: Student[]     | Actual: Student[]',
        '-> Compile-time: IEnumerable`1 | Actual: Student[]',
        '-> Compile-time: IQueryable`1  | Actual: EnumerableQuery`1'
    ], 'Program.cs')
    add_bottom_banner(s97, "📌 Застосування AsEnumerable в EF Core:",
                      "Використовуйте AsEnumerable(), коли метод не може бути трансльований у SQL і його потрібно виконати на боці клієнта.")

    # ==========================================
    # SLIDE 98: CAST<T>
    # ==========================================
    s98 = prs.slides.add_slide(blank_layout)
    add_bg(s98)
    add_header(s98, "16 · Перетворення (Conversion)", "Оператор приведення Cast<T>: Робота з нетипізованими колекціями", 98)
    add_theory_card(s98, 0.8, 1.55, 5.72, 4.25,
        "Приведення типів елементів у послідовності", ACCENT_INDIGO,
        [
            ("Призначення Cast<T>:", "Перетворює непараметризовану колекцію (IEnumerable з .NET 1.1, наприклад ArrayList) на типізовану IEnumerable<T>."),
            ("Приведення кожного елемента:", "Виконує явне приведення (T)element для кожного запису послідовності."),
            ("Ризик винятку:", "Якщо хоча б один елемент у колекції має несумісний тип, виникає виняток InvalidCastException."),
            ("Безпечна альтернатива:", "Якщо колекція містить різнорідні об'єкти, використовуйте OfType<T>(), який просто ігнорує несумісні типи.")
        ],
        "⚡ Результат приведення:",
        "Отримуємо строго типізований IEnumerable<Student>, з яким можна працювати за допомогою сучасного LINQ."
    )
    add_code_panel(s98, 6.81, 1.55, 5.72, 4.25, [
        "// Нетипізована колекція legacy .NET 1.1",
        "System.Collections.ArrayList list = new System.Collections.ArrayList();",
        "list.Add(new Student() { StudentID = 1, StudentName = \"John\", Age = 18 });",
        "list.Add(new Student() { StudentID = 2, StudentName = \"Steve\", Age = 21 });",
        "list.Add(new Student() { StudentID = 3, StudentName = \"Bill\", Age = 25 });",
        "",
        "// Приведення до IEnumerable<Student>",
        "var studentEnumerable = list.Cast<Student>();",
        "",
        "foreach (Student std in studentEnumerable)",
        "    Console.WriteLine(std.StudentName);",
        "",
        "-> Output:",
        "-> John",
        "-> Steve",
        "-> Bill"
    ], "Program.cs")
    add_bottom_banner(s98, "💡 Відмінність від OfType:",
                      "Cast вимагає строгої відповідності 100% елементів типу T, тоді як OfType безпечно фільтрує невідповідні елементи без винятків.")

    # ==========================================
    # SLIDE 99: TOARRAY ТА TOLIST
    # ==========================================
    s99 = prs.slides.add_slide(blank_layout)
    add_bg(s99)
    add_header(s99, "16 · Перетворення (Conversion)", "Оператори матеріалізації: ToArray та ToList", 99)
    add_theory_card(s99, 0.8, 1.55, 5.72, 4.25,
        "Фіксація результатів запиту в оперативній пам'яті", ACCENT_INDIGO,
        [
            ("Негайне виконання:", "ToArray та ToList примусово виконують запит у момент їх виклику (Immediate Execution)."),
            ("ToArray():", "Створює одновимірний масив фіксованої довжини T[]. Ідеально для передачі в API з мінімальним оверхедом."),
            ("ToList():", "Створює динамічний список List<T>, що дозволяє додавати та видаляти елементи."),
            ("Кешування результату:", "Захищає від повторного виконання дорогих SQL-запитів при багаторазовому читанні даних.")
        ],
        "⚡ Результат перетворення:",
        "Отримуємо масив string[] або список List<string> з повним набором даних у пам'яті."
    )
    add_code_panel(s99, 6.81, 1.55, 5.72, 4.25, [
        "IList<string> strList = new List<string>() {",
        "    \"One\", \"Two\", \"Three\", \"Four\", \"Three\"",
        "};",
        "",
        "// Перетворення списку в одновимірний масив",
        "string[] strArray = strList.ToArray();",
        "",
        "// Перетворення масиву назад у типізований список",
        "IList<string> list = strArray.ToList();",
        "",
        "Console.WriteLine(\"Масив: \" + strArray.Length + \" ел.\");",
        "Console.WriteLine(\"Список: \" + list.Count + \" ел.\");",
        "",
        "-> Output:",
        "-> Масив: 5 ел.",
        "-> Список: 5 ел."
    ], "Program.cs")
    add_bottom_banner(s99, "📌 Золоте правило:",
                      "Завжди матеріалізуйте дані через .ToList() перед закриттям з'єднання з базою даних або передачею моделі у представлення.")

    # ==========================================
    # SLIDE 100: TODICTIONARY
    # ==========================================
    s100 = prs.slides.add_slide(blank_layout)
    add_bg(s100)
    add_header(s100, "16 · Перетворення (Conversion)", "Оператор ToDictionary: Створення асоціативної таблиці за ключем", 100)
    add_theory_card(s100, 0.8, 1.55, 5.72, 4.25,
        "Матеріалізація в Dictionary<TKey, TValue>", ACCENT_INDIGO,
        [
            ("Селектор ключа:", "ToDictionary вимагає обов'язкової функції вибору ключа, наприклад s => s.StudentID."),
            ("Швидкий доступ O(1):", "Словник забезпечує миттєвий пошук студента за його ID без необхідності перебору колекції."),
            ("Вимога унікальності:", "Усі ключі повинні бути суворо унікальними! Якщо трапиться два однакових StudentID, метод викине ArgumentException."),
            ("Селектор значення:", "Можна опційно передати другу лямбду для збереження лише окремого поля: s => s.StudentName.")
        ],
        "⚡ Результат матеріалізації:",
        "Створено Dictionary<int, Student>, де ключем є StudentID, а значенням — об'єкт студента."
    )
    add_code_panel(s100, 6.81, 1.55, 5.72, 4.25, [
        "IList<Student> studentList = new List<Student>() {",
        "    new Student() { StudentID = 1, StudentName = \"John\", Age = 18 },",
        "    new Student() { StudentID = 2, StudentName = \"Steve\", Age = 21 },",
        "    new Student() { StudentID = 3, StudentName = \"Bill\", Age = 18 },",
        "    new Student() { StudentID = 4, StudentName = \"Ram\", Age = 20 },",
        "    new Student() { StudentID = 5, StudentName = \"Ron\", Age = 21 }",
        "};",
        "",
        "// Конвертація в Dictionary з ключем StudentID",
        "IDictionary<int, Student> studentDict = studentList",
        "    .ToDictionary<Student, int>(s => s.StudentID);",
        "",
        "foreach (var key in studentDict.Keys)",
        "    Console.WriteLine($\"Key: {key}, Value: {studentDict[key].StudentName}\");",
        "",
        "-> Key: 1, Value: John",
        "-> Key: 2, Value: Steve",
        "-> Key: 3, Value: Bill..."
    ], "Program.cs")
    add_bottom_banner(s100, "💡 Що робити з дублікатами ключів:",
                      "Якщо ключі можуть повторюватися, замість ToDictionary використовуйте ToLookup() або GroupBy(), які підтримують групи 1:N.")

    # ==========================================
    # SLIDE 101: DEFERRED EXECUTION ВСТУП
    # ==========================================
    s101 = prs.slides.add_slide(blank_layout)
    add_bg(s101)
    add_header(s101, "17 · Моделі виконання (Execution Models)", "Відкладене виконання (Deferred Execution) у LINQ", 101)
    add_theory_card(s101, 0.8, 1.55, 5.72, 4.25,
        "Концепція відкладеного обчислення", ACCENT_INDIGO,
        [
            ("Що таке Deferred Execution:", "Оголошення LINQ-запиту НЕ виконує сам запит і не перебирає елементи колекції в пам'яті."),
            ("Збереження плану запиту:", "Змінна запиту teenAgerStudents зберігає лише команду та логіку обробки (структуру IEnumerable)."),
            ("Точка виконання:", "Запит реально виконується тільки тоді, коли дані безпосередньо запитуються (наприклад, у циклі foreach)."),
            ("Оптимізація ресурсів:", "Дозволяє уникати непотрібних проміжних обчислень та об'єднувати фільтри перед відправкою до бази даних.")
        ],
        "⚡ Дві критичні точки в коді:",
        "1. Оголошення запиту: нуль операцій з елементами. 2. Цикл foreach: послідовне виконання ітератора."
    )
    add_code_panel(s101, 6.81, 1.55, 5.72, 4.25, [
        "IList<Student> studentList = new List<Student>() {",
        "    new Student() { StudentID = 1, StudentName = \"John\", Age = 18 },",
        "    new Student() { StudentID = 2, StudentName = \"Steve\", Age = 21 },",
        "    new Student() { StudentID = 3, StudentName = \"Bill\", Age = 18 },",
        "    new Student() { StudentID = 4, StudentName = \"Ram\", Age = 20 },",
        "    new Student() { StudentID = 5, StudentName = \"Ron\", Age = 21 }",
        "};",
        "",
        "// ТОЧКА 1: Запит ТУТ НЕ ВИКОНУЄТЬСЯ!",
        "var teenAgerStudents = from s in studentList",
        "                       where s.Age > 12 && s.Age < 20",
        "                       select s;",
        "",
        "// ТОЧКА 2: Запит РЕАЛЬНО ВИКОНУЄТЬСЯ ТУТ!",
        "foreach (Student teenStudent in teenAgerStudents)",
        "    Console.WriteLine(\"Name: \" + teenStudent.StudentName);",
        "",
        "-> Output:",
        "-> Name: John",
        "-> Name: Bill"
    ], "Program.cs")
    add_bottom_banner(s101, "💡 Універсальність:",
                      "Deferred execution діє як для локальних колекцій у пам'яті (LINQ to Objects), так і для віддалених джерел (LINQ to Entities, SQL, XML).")

    # ==========================================
    # SLIDE 102: QUERY RE-EVALUATION
    # ==========================================
    s102 = prs.slides.add_slide(blank_layout)
    add_bg(s102)
    add_header(s102, "17 · Моделі виконання (Execution Models)", "Повторне обчислення запиту (Query Re-evaluation)", 102)
    add_theory_card(s102, 0.8, 1.55, 5.72, 4.25,
        'Динамічний доступ до актуальних даних', ACCENT_INDIGO,
        [
            ('Відсутність кешу за замовчуванням:', "Оскільки відкладений запит не зберігає результатів у пам'яті, кожен новий цикл foreach звертається до джерела наново."),
            ('Модифікація колекції:', 'Якщо після першого виконання додати нового студента (Martin, 16 років), він автоматично врахується при повторному обході!'),
            ('Один запит — актуальні дані:', 'Запит можна оголосити один раз і повторно використовувати його для отримання найсвіжіших даних з джерела.'),
            ('Уникнення подвійного обчислення:', 'Якщо джерелом є база даних або складний розрахунок, багаторазовий перехід запиту виконує повторні мережеві запити. Фіксуйте результат через .ToList().')
        ],
        '⚡ Результати двох циклів:',
        'Перший foreach: John, Bill. Другий foreach (після Add): John, Bill, Martin!'
    )
    add_code_panel(s102, 6.81, 1.55, 5.72, 4.25, [
        '// Перше виконання запиту:',
        'foreach (Student teenStudent in teenAgerStudents)',
        '    Console.WriteLine("Name: " + teenStudent.StudentName);',
        '// -> John, Bill',
        '',
        '// Додаємо нового студента у вихідний список:',
        'studentList.Add(new Student() {',
        '    StudentID = 6, StudentName = "Martin", Age = 16',
        '});',
        '',
        '// Друге виконання того ж самого запиту повертає свіжі дані:',
        'foreach (Student teenStudent in teenAgerStudents)',
        '    Console.WriteLine("Name: " + teenStudent.StudentName);',
        '',
        '-> Output другого циклу:',
        '-> Name: John',
        '-> Name: Bill',
        '-> Name: Martin (доданий елемент враховано!)'
    ], 'Program.cs')
    add_bottom_banner(s102, "📌 Важливий побічний ефект:",
                      "Якщо колекція змінюється між ітераціями, результати будуть різними. Якщо потрібен незмінний знімок даних — використовуйте .ToList().")

    # ==========================================
    # SLIDE 103: РЕАЛІЗАЦІЯ ДЕФЕРЕД ЕКЗЕКЬЮШН (YIELD RETURN)
    # ==========================================
    s103 = prs.slides.add_slide(blank_layout)
    add_bg(s103)
    add_header(s103, "17 · Моделі виконання (Execution Models)", "Реалізація відкладеного виконання: Механізм yield return", 103)
    add_theory_card(s103, 0.8, 1.55, 5.72, 4.25,
        'Як побудовані ітератори BCL зсередини', ACCENT_INDIGO,
        [
            ('Ключове слово yield return:', 'Перетворює звичайний C# метод на генератор ітератора, що реалізує інтерфейс IEnumerator<T>.'),
            ('Скінченний автомат:', 'Компілятор генерує прихований клас State Machine, який зберігає стан виконання між викликами MoveNext().'),
            ('Потокова обробка:', 'Метод не повертає всю колекцію відразу — він повертає рівно один елемент і призупиняється до наступного запиту.'),
            ('Очищення ресурсів (Dispose):', 'Ітератор гарантує виклик блоків finally та Dispose, коли цикл foreach завершується або переривається за допомогою break.')
        ],
        '⚡ Власний LINQ оператор:',
        'Метод GetTeenAgerStudents фільтрує студентів потоково, логуючи кожне звернення до елемента.'
    )
    add_code_panel(s103, 6.81, 1.55, 5.72, 4.25, [
        'public static class EnumerableExtensionMethods {',
        '    // Кастомний метод-розширення з відкладеним виконанням',
        '    public static IEnumerable<Student> GetTeenAgerStudents(',
        '        this IEnumerable<Student> source)',
        '    {',
        '        foreach (Student std in source)',
        '        {',
        '            Console.WriteLine("Accessing student {0}", std.StudentName);',
        '',
        '            if (std.Age > 12 && std.Age < 20)',
        '            {',
        '                yield return std; // пауза виконання до наступного MoveNext()',
        '            }',
        '        }',
        '    }',
        '}'
    ], 'EnumerableExtensionMethods.cs')
    add_bottom_banner(s103, "💡 Магія yield return:",
                      "Код методу взагалі не виконується в момент виклику GetTeenAgerStudents() — він стартує тільки при першому зверненні до ітератора!")

    # ==========================================
    # SLIDE 104: ТРАСУВАННЯ YIELD RETURN
    # ==========================================
    s104 = prs.slides.add_slide(blank_layout)
    add_bg(s104)
    add_header(s104, "17 · Моделі виконання (Execution Models)", "Трасування роботи ітератора з yield return на практиці", 104)
    add_theory_card(s104, 0.8, 1.55, 5.72, 4.25,
        'Доведення відкладеності у виводі консолі', ACCENT_INDIGO,
        [
            ('Крок 1 (Оголошення):', 'Виклик GetTeenAgerStudents() у секції from s in ... не генерує ЖОДНОГО повідомлення в консолі!'),
            ('Крок 2 (Ітерація):', 'У циклі foreach починає викликатися MoveNext(). Метод виконується поетапно для кожного студента.'),
            ('Аналіз виводу:', "Рядки 'Accessing student...' чергуються з виведенням імені в циклі 'Name: ...'."),
            ('Ранній вихід:', 'Якби в циклі стояв break після першого студента, перевірка решти списку навіть не розпочалася б.')
        ],
        '⚡ Результат трасування:',
        'Повний синхронний контроль над споживанням даних без виділення проміжних масивів.'
    )
    add_code_panel(s104, 6.81, 1.55, 5.72, 4.25, [
        '// 1. Оголошення запиту - GetTeenAgerStudents НЕ викликається:',
        'var teenAgerStudents = from s in studentList.GetTeenAgerStudents()',
        '                       select s;',
        '',
        'Console.WriteLine("Запит оголошено. Починаємо foreach:");',
        '',
        '// 2. Ітерація - ТІЛЬКИ ТУТ починається обробка елементів:',
        'foreach (Student teenStudent in teenAgerStudents)',
        '{',
        '    Console.WriteLine("-> Знайдено: " + teenStudent.StudentName);',
        '}',
        '',
        '-> Console Trace:',
        '-> Accessing student John -> Знайдено: John',
        '-> Accessing student Steve',
        '-> Accessing student Bill -> Знайдено: Bill',
        '-> Accessing student Ram',
        '-> Accessing student Ron'
    ], 'Program.cs')
    add_bottom_banner(s104, "📌 Підсумок механізму:",
                      "Це фундаментальна основа функціонального програмування в C#: потокова конвеєризація даних (Pipeline Processing).")

    # ==========================================
    # SLIDE 105: IMMEDIATE EXECUTION
    # ==========================================
    s105 = prs.slides.add_slide(blank_layout)
    add_bg(s105)
    add_header(s105, "17 · Моделі виконання (Execution Models)", "Негайне виконання запитів (Immediate Execution)", 105)
    add_theory_card(s105, 0.8, 1.55, 5.72, 4.25,
        'Примусове обчислення в точці виклику', ACCENT_INDIGO,
        [
            ('Що таке Immediate Execution:', 'Повна протилежність відкладеному виконанню. Змушує запит негайно обробити всі дані і зберегти знімок результату.'),
            ('Оператори негайного виконання:', "Всі 'To' оператори (ToList, ToArray, ToDictionary, ToLookup), а також агрегатні функції (Count, Average, First, Single, Max)."),
            ('Ізоляція від змін:', 'Результат зберігається у створеній колекції (List<Student>), тому подальші модифікації studentList уже не впливають на збережені дані.'),
            ("Використання пам'яті:", 'Матеріалізація завантажує всю вибірку в RAM, тому для великих обсягів даних слід обмежувати вибірку через Take().')
        ],
        '⚡ Приклад у Method Syntax:',
        "Метод .ToList() негайно фільтрує список і повертає готовий об'єкт List<Student>."
    )
    add_code_panel(s105, 6.81, 1.55, 5.72, 4.25, [
        '// Негайне виконання за допомогою .ToList():',
        'IList<Student> teenAgerStudents = studentList',
        '    .Where(s => s.Age > 12 && s.Age < 20)',
        '    .ToList();',
        '',
        '// Запит повністю виконано прямо зараз у точці .ToList()!',
        "// Змінна teenAgerStudents вже містить 2 об'єкти в пам'яті.",
        'Console.WriteLine($"Завантажено студентів: {teenAgerStudents.Count}");',
        '',
        'foreach (Student teenStudent in teenAgerStudents)',
        '{',
        '    Console.WriteLine($"ID: {teenStudent.StudentID}, Name: {teenStudent.StudentName}");',
        '}',
        '',
        '-> Output:',
        '-> Завантажено студентів: 2',
        '-> ID: 1, Name: John',
        '-> ID: 3, Name: Bill'
    ], 'Program.cs')
    add_bottom_banner(s105, "💡 Коли застосовувати Immediate Execution:",
                      "Завжди, коли потрібно зафіксувати дані в пам'яті (кеш), від'єднатися від БД або запобігти повторним важким обрахункам.")

    # ==========================================
    # SLIDE 106: IMMEDIATE EXECUTION В QUERY SYNTAX
    # ==========================================
    s106 = prs.slides.add_slide(blank_layout)
    add_bg(s106)
    add_header(s106, "17 · Моделі виконання (Execution Models)", "Негайне виконання у синтаксисі запиту (Query Syntax)", 106)
    add_theory_card(s106, 0.8, 1.55, 5.72, 4.25,
        'Керування виконанням через круглі дужки', ACCENT_INDIGO,
        [
            ('Природа Query Syntax:', 'Синтаксис запиту (from ... select) за замовчуванням ЗАВЖДИ є відкладеним (повертає IEnumerable).'),
            ('Прийом обгортання:', 'Щоб виконати такий запит негайно, весь вираз береться в круглі дужки і викликається оператор матеріалізації: (...) .ToList().'),
            ('Перегляд у Debugger:', 'У відкладеному запиті налагоджувач показує внутрішній тип ітератора. У матеріалізованому — одразу відображає Count = 2 та масив елементів.'),
            ('Матеріалізація скалярів:', 'Будь-який запит Query Syntax можна також обгорнути в дужки для виклику First(), Count(), Any() або ToDictionary().')
        ],
        '⚡ Практична різниця:',
        "Без .ToList(): запит чекає ітерації. З .ToList(): у пам'яті одразу створено готовий список із 2 студентів."
    )
    add_code_panel(s106, 6.81, 1.55, 5.72, 4.25, [
        '// 1. Відкладений синтаксис запиту:',
        'var deferredQuery = from s in studentList',
        '                    where s.Age > 12 && s.Age < 20',
        '                    select s;',
        '',
        '// 2. Негайне виконання через круглі дужки:',
        'IList<Student> teenAgerStudents = (',
        '    from s in studentList',
        '    where s.Age > 12 && s.Age < 20',
        '    select s',
        ').ToList();',
        '',
        '// QuickWatch / Debugger View:',
        '// teenAgerStudents.Count = 2',
        '// [0] = {StudentID = 1, StudentName = "John"}',
        '// [1] = {StudentID = 3, StudentName = "Bill"}'
    ], 'Program.cs')
    add_bottom_banner(s106, "📌 Патерн для команди:",
                      "Обгортання запиту в дужки з викликом .ToList() або .ToArray() є загальноприйнятим стандартом у комерційній розробці на C#.")

    # ==========================================
    # SLIDE 107: LET KEYWORD
    # ==========================================
    s107 = prs.slides.add_slide(blank_layout)
    add_bg(s107)
    add_header(s107, "18 · Ключові слова запитів (Query Keywords)", "Ключове слово let: Оголошення проміжних змінних у запиті", 107)
    add_theory_card(s107, 0.8, 1.55, 5.72, 4.25,
        "Оптимізація та читабельність LINQ запитів", ACCENT_INDIGO,
        [
            ("Проблема без let:", "Якщо потрібно використовувати один і той самий вираз (наприклад, s.StudentName.ToLower()) у where та select, його доводиться писати двічі."),
            ("Призначення let:", "Ключове слово let вводить нову змінну в запиті та зберігає в ній проміжний результат обчислення."),
            ("Прозорий ідентифікатор:", "Компілятор C# створює анонімний тип (Transparent Identifier), що містить вихідний об'єкт та нову змінну."),
            ("Підвищення продуктивності:", "Метод ToLower() викликається лише один раз для кожного запису замість двох чи більше разів.")
        ],
        "⚡ Результат виконання:",
        "Вибірка містить імена в нижньому регістрі: 'ram', 'ron'."
    )
    add_code_panel(s107, 6.81, 1.55, 5.72, 4.25, [
        "// БЕЗ LET: ToLower() викликається двічі для кожного студента",
        "var badQuery = from s in studentList",
        "               where s.StudentName.ToLower().StartsWith(\"r\")",
        "               select s.StudentName.ToLower();",
        "",
        "// З LET: обчислюємо один раз і повторно використовуємо",
        "var lowercaseStudentNames = from s in studentList",
        "    let lowercaseName = s.StudentName.ToLower()",
        "    where lowercaseName.StartsWith(\"r\")",
        "    select lowercaseName;",
        "",
        "foreach (var name in lowercaseStudentNames)",
        "    Console.WriteLine(name);",
        "",
        "-> Output:",
        "-> ram",
        "-> ron"
    ], "Program.cs")
    add_bottom_banner(s107, "💡 Best Practice:",
                      "Використовуйте let для збереження складних математичних формул, викликів методів та розбору рядків усередині запиту.")

    # ==========================================
    # SLIDE 108: INTO KEYWORD
    # ==========================================
    s108 = prs.slides.add_slide(blank_layout)
    add_bg(s108)
    add_header(s108, "18 · Ключові слова запитів (Query Keywords)", "Ключове слово into: Продовження запиту (Query Continuation)", 108)
    add_theory_card(s108, 0.8, 1.55, 5.72, 4.25,
        'Конвеєризація запиту після select або group', ACCENT_INDIGO,
        [
            ('Призначення into:', 'Дозволяє продовжити виконання запиту після секцій select, group by або join, створюючи новий ідентифікатор діапазону.'),
            ('Область видимості:', 'Увага! Після ключового слова into попередня змінна діапазону (s) виходить з області видимості і більше недоступна!'),
            ('Багаторівнева фільтрація:', "У наведеному прикладі ми спочатку вибираємо підлітків, а потім продовжуємо фільтрацію за першою літерою імені ('B')."),
            ('Завершення курсу LINQ:', 'Повне поєднання Query та Method синтаксису дає найпотужніший інструмент роботи з даними у .NET 10.')
        ],
        '⚡ Фінальний результат запиту:',
        'Повертається студент Bill, який пройшов обидва етапи фільтрації.'
    )
    add_code_panel(s108, 6.81, 1.55, 5.72, 4.25, [
        "// Продовження запиту після select за допомогою 'into'",
        'var teenAgerStudents = ',
        '    from s in studentList',
        '    where s.Age > 12 && s.Age < 20',
        '    select s',
        "    into teenStudents // 's' більше недоступна, працюємо з teenStudents",
        '    where teenStudents.StudentName.StartsWith("B")',
        '    select teenStudents;',
        '',
        'Console.WriteLine("--- Студенти після into-фільтрації ---");',
        'foreach (var teen in teenAgerStudents)',
        '{',
        '    Console.WriteLine($"ID: {teen.StudentID}, Name: {teen.StudentName}, Age: {teen.Age}");',
        '}',
        '',
        '-> Output:',
        '-> ID: 3, Name: Bill, Age: 18'
    ], 'Program.cs')
    add_bottom_banner(s108, "🎓 Вітаємо із завершенням курсу!",
                      "Ви опанували всі 18 ключових тем LINQ: від базової архітектури та лямбд до складних реляційних з'єднань, секціонування та оптимізації виконання!")

