import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==========================================
# ENTERPRISE COLOR PALETTE
# ==========================================
BG_STUDIO = RGBColor(248, 250, 252)       # Slate 50
CARD_BG = RGBColor(255, 255, 255)         # Pure White
CARD_BORDER = RGBColor(203, 213, 225)     # Slate 300
TEXT_MAIN = RGBColor(15, 23, 42)          # Slate 900
TEXT_MUTED = RGBColor(51, 65, 85)         # Slate 700
TEXT_SUBTLE = RGBColor(100, 116, 139)     # Slate 500

# ACCENTS
ACCENT_INDIGO = RGBColor(79, 70, 229)     # Indigo 600
ACCENT_SKY = RGBColor(2, 132, 199)        # Sky 600
ACCENT_EMERALD = RGBColor(5, 150, 105)    # Emerald 600
ACCENT_VIOLET = RGBColor(124, 58, 237)    # Violet 600
ACCENT_FUCHSIA = RGBColor(192, 38, 211)   # Fuchsia 600
ACCENT_AMBER = RGBColor(217, 119, 6)      # Amber 600
ACCENT_TEAL = RGBColor(13, 148, 136)      # Teal 600
ACCENT_CRIMSON = RGBColor(220, 38, 38)    # Red 600

# DARK CODE THEME
CODE_BG = RGBColor(15, 23, 42)            # Dark Slate 900
CODE_BORDER = RGBColor(51, 65, 85)        # Slate 700
CODE_TEXT = RGBColor(226, 232, 240)       # Light Slate 200
CODE_ACCENT = RGBColor(56, 189, 248)      # Sky 400 (Keywords)
CODE_COMMENT = RGBColor(52, 211, 153)     # Emerald 400 (Comments)
CODE_OUTPUT = RGBColor(251, 191, 36)      # Amber 400 (Output/Console)

def add_bg(slide, bg_col=BG_STUDIO):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.500))
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_col
    bg.line.fill.background()
    return bg

def add_header(slide, module_tag, title_text, slide_num, total_slides=108, accent_rgb=ACCENT_INDIGO):
    # Tag rectangle (sharp rectangle)
    tag_w = max(3.4, len(module_tag) * 0.088 + 0.6)
    tag_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.38), Inches(tag_w), Inches(0.32))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = accent_rgb
    tag_box.line.fill.background()
    tag_box.text_frame.margin_left = Inches(0.08)
    tag_box.text_frame.margin_right = Inches(0.08)
    tag_box.text_frame.margin_top = Inches(0.03)
    tag_box.text_frame.margin_bottom = Inches(0.03)
    p = tag_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = module_tag.upper()
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.74), Inches(11.73), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Calibri"
    p_title.font.size = Pt(23)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_MAIN

    # Footer Left
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(6.5), Inches(0.28))
    p_foot = footer_box.text_frame.paragraphs[0]
    p_foot.text = "ООП · ЛЕКЦІЯ №10 · LINQ У C# ТА .NET 10"
    p_foot.font.name = "Calibri"
    p_foot.font.size = Pt(9.5)
    p_foot.font.color.rgb = TEXT_SUBTLE

    # Footer Right
    num_box = slide.shapes.add_textbox(Inches(11.33), Inches(7.12), Inches(1.2), Inches(0.28))
    p_num = num_box.text_frame.paragraphs[0]
    p_num.alignment = PP_ALIGN.RIGHT
    p_num.text = f"{slide_num:03d} / {total_slides:03d}"
    p_num.font.name = "Calibri"
    p_num.font.size = Pt(9.5)
    p_num.font.bold = True
    p_num.font.color.rgb = TEXT_SUBTLE

def add_bottom_banner(slide, label, text, banner_type="tip"):
    top_y = 5.95
    w = 11.733
    h = 1.05
    bg_col = RGBColor(238, 242, 255) if banner_type == "tip" else RGBColor(254, 243, 199)
    border_col = ACCENT_INDIGO if banner_type == "tip" else ACCENT_AMBER
    lbl_col = ACCENT_INDIGO if banner_type == "tip" else RGBColor(180, 83, 9)

    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_y), Inches(w), Inches(h))
    banner.fill.solid()
    banner.fill.fore_color.rgb = bg_col
    banner.line.color.rgb = border_col
    banner.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(Inches(0.95), Inches(top_y + 0.08), Inches(w - 0.3), Inches(h - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p1 = tf.paragraphs[0]
    p1.text = label
    p1.font.name = "Calibri"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = lbl_col

    p2 = tf.add_paragraph()
    p2.text = text
    p2.font.name = "Calibri"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MAIN

def add_code_panel(slide, left, top, width, height, lines, ref_file=None):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = CODE_BORDER
    box.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Inches(left + 0.12), Inches(top + 0.08), Inches(width - 0.24), Inches(height - 0.16))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)

    num_lines = len(lines)
    if height > 3.0:
        if num_lines >= 22:
            code_font_size = Pt(9.5)
            p_line_sp = 1.02
            p_space_after = Pt(0)
        elif num_lines >= 17:
            code_font_size = Pt(10)
            p_line_sp = 1.08
            p_space_after = Pt(0.5)
        elif num_lines >= 13:
            code_font_size = Pt(10.5)
            p_line_sp = 1.15
            p_space_after = Pt(1.5)
        else:
            code_font_size = Pt(11)
            p_line_sp = 1.25
            p_space_after = Pt(2.5)
    else:
        if num_lines >= 9:
            code_font_size = Pt(8.5)
            p_line_sp = 1.02
            p_space_after = Pt(0)
        elif num_lines >= 7:
            code_font_size = Pt(9.5)
            p_line_sp = 1.08
            p_space_after = Pt(1)
        else:
            code_font_size = Pt(10)
            p_line_sp = 1.18
            p_space_after = Pt(2)

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = p_line_sp
        p.space_after = p_space_after
        run = p.add_run()
        run.text = line
        run.font.name = "Consolas"
        run.font.size = code_font_size

        sline = line.strip()
        if sline.startswith("//"):
            run.font.color.rgb = CODE_COMMENT
        elif sline.startswith("->") or sline.startswith("Output:") or sline.startswith("Key :") or sline.startswith("1st Element"):
            run.font.color.rgb = CODE_OUTPUT
        elif any(sline.startswith(kw) for kw in ["var ", "from ", "where ", "select ", "group ", "into ", "order", "join ", "public ", "static ", "new ", "foreach", "class "]):
            run.font.color.rgb = CODE_ACCENT
        else:
            run.font.color.rgb = CODE_TEXT

    if ref_file:
        p_ref = tf.add_paragraph()
        p_ref.space_before = Pt(4)
        r_ref = p_ref.add_run()
        r_ref.text = f"📍 {ref_file}"
        r_ref.font.name = "Consolas"
        r_ref.font.size = Pt(9.5)
        r_ref.font.color.rgb = CODE_ACCENT

def add_theory_card(slide, left, top, width, height, title, title_color, bullets, sub_title=None, sub_text=None, sub_accent=ACCENT_INDIGO):
    c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = CARD_BORDER
    c.line.width = Pt(1)

    has_many_bullets = len(bullets) >= 4
    title_font_sz = Pt(17.0) if has_many_bullets else Pt(18.0)
    b_head_sz = Pt(12) if has_many_bullets else Pt(13)
    b_text_sz = Pt(11.5) if has_many_bullets else Pt(12.5)
    sp_before = Pt(3.5) if has_many_bullets else Pt(6.5)
    line_sp = 1.10 if has_many_bullets else 1.20

    sub_h = 1.10 if (sub_title and sub_text) else 0
    tb_h = (height - sub_h - 0.25) if (sub_title and sub_text) else (height - 0.16)
    tb = slide.shapes.add_textbox(Inches(left + 0.12), Inches(top + 0.08), Inches(width - 0.24), Inches(tb_h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)

    p_head = tf.paragraphs[0]
    p_head.text = title
    p_head.font.name = "Calibri"
    p_head.font.size = title_font_sz
    p_head.font.bold = True
    p_head.font.color.rgb = title_color

    for b_title, b_text in bullets:
        p = tf.add_paragraph()
        p.space_before = sp_before
        p.line_spacing = line_sp
        r1 = p.add_run()
        r1.text = f"• {b_title} "
        r1.font.name = "Calibri"
        r1.font.size = b_head_sz
        r1.font.bold = True
        r1.font.color.rgb = TEXT_MAIN
        r2 = p.add_run()
        r2.text = b_text
        r2.font.name = "Calibri"
        r2.font.size = b_text_sz
        r2.font.color.rgb = TEXT_MUTED

    if sub_title and sub_text:
        sub_y = top + height - sub_h - 0.10
        sub = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.12), Inches(sub_y), Inches(width - 0.24), Inches(sub_h))
        sub.fill.solid()
        sub.fill.fore_color.rgb = RGBColor(241, 245, 249)
        sub.line.color.rgb = RGBColor(203, 213, 225)
        sub.line.width = Pt(1)
        tf_sub = sub.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = Inches(0.08)
        tf_sub.margin_right = Inches(0.08)
        tf_sub.margin_top = Inches(0.05)
        tf_sub.margin_bottom = Inches(0.05)
        p_s1 = tf_sub.paragraphs[0]
        p_s1.text = sub_title
        p_s1.font.name = "Calibri"
        p_s1.font.size = Pt(12)
        p_s1.font.bold = True
        p_s1.font.color.rgb = sub_accent
        p_s2 = tf_sub.add_paragraph()
        p_s2.space_before = Pt(2)
        p_s2.line_spacing = 1.15
        p_s2.text = sub_text
        p_s2.font.name = "Calibri"
        p_s2.font.size = Pt(11.5)
        p_s2.font.color.rgb = TEXT_MAIN

def add_dual_comparison_card(slide, left, top, width, height, title, title_color, bullets, code_lines):
    c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = CARD_BORDER
    c.line.width = Pt(1)

    tb = slide.shapes.add_textbox(Inches(left + 0.12), Inches(top + 0.10), Inches(width - 0.24), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p_head = tf.paragraphs[0]
    p_head.text = title
    p_head.font.name = "Calibri"
    p_head.font.size = Pt(17)
    p_head.font.bold = True
    p_head.font.color.rgb = title_color

    for b_title, b_text in bullets:
        p = tf.add_paragraph()
        p.space_before = Pt(3)
        r1 = p.add_run()
        r1.text = f"• {b_title} "
        r1.font.name = "Calibri"
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = TEXT_MAIN
        r2 = p.add_run()
        r2.text = b_text
        r2.font.name = "Calibri"
        r2.font.size = Pt(11.5)
        r2.font.color.rgb = TEXT_MUTED

    add_code_panel(slide, left + 0.12, top + 2.30, width - 0.24, 1.80, code_lines)

def add_table_card(slide, left, top, width, height, headers, rows, col_widths, accent_color=ACCENT_INDIGO):
    # Outer container
    c = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = CARD_BORDER
    c.line.width = Pt(1)

    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(left + 0.1), Inches(top + 0.1), Inches(width - 0.2), Inches(height - 0.2))
    table = table_shape.table

    for i, w in enumerate(col_widths):
        table.columns[i].width = Inches(w)

    # Header Row
    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = accent_color
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = h_text
        run.font.name = "Calibri"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(rows):
        is_even = (row_idx % 2 == 0)
        row_bg = RGBColor(255, 255, 255) if is_even else RGBColor(248, 250, 252)
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = row_bg
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(cell_value)
            run.font.name = "Consolas" if ("(" in str(cell_value) or "=>" in str(cell_value) or col_idx == 0) else "Calibri"
            run.font.size = Pt(10)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = ACCENT_INDIGO
            else:
                run.font.color.rgb = TEXT_MAIN
